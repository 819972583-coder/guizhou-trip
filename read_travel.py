import sys, os, zipfile, re
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

XLSX = r"C:\Users\SYP\Desktop\【6天5晚】2026无限黔徒-第26期.xlsx"
DOCX = r"C:\Users\SYP\Downloads\贵州徒步每日穿衣推荐_1789801280986177022.docx"

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

print("=" * 80)
print("XLSX CONTENT")
print("=" * 80)
wb = load_workbook(XLSX, data_only=True)
for ws in wb.worksheets:
    print("\n" + "#" * 70)
    print(f"# SHEET: {ws.title}  (dims={ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column})")
    print("#" * 70)
    for row in ws.iter_rows(values_only=True):
        cells = []
        for c in row:
            if c is None:
                cells.append("")
            else:
                cells.append(str(c))
        line = " | ".join(cells).rstrip(" |")
        if line.strip():
            print(line)

print("\n" + "=" * 80)
print("DOCX CONTENT")
print("=" * 80)

def docx_text(path):
    out = []
    with zipfile.ZipFile(path) as z:
        # document.xml
        names = z.namelist()
        doc_name = "word/document.xml"
        if doc_name in names:
            data = z.read(doc_name)
            root = ET.fromstring(data)
            body = root.find(f"{W_NS}body")
            def walk(elem):
                for child in elem:
                    tag = child.tag
                    if tag == f"{W_NS}p":
                        # paragraph
                        txt = "".join(t.text or "" for t in child.iter(f"{W_NS}t"))
                        # detect style
                        style = ""
                        pPr = child.find(f"{W_NS}pPr")
                        if pPr is not None:
                            pStyle = pPr.find(f"{W_NS}pStyle")
                            if pStyle is not None:
                                style = pStyle.get(f"{W_NS}val", "")
                        out.append(("P", style, txt))
                    elif tag == f"{W_NS}tbl":
                        rows = []
                        for tr in child.findall(f"{W_NS}tr"):
                            cells = []
                            for tc in tr.findall(f"{W_NS}tc"):
                                ctext = "".join(t.text or "" for t in tc.iter(f"{W_NS}t"))
                                cells.append(ctext)
                            rows.append(" || ".join(cells))
                        out.append(("TABLE", "", "\n".join(rows)))
                    else:
                        walk(child)
            if body is not None:
                walk(body)
    return out

for kind, style, txt in docx_text(DOCX):
    if kind == "TABLE":
        print("\n--- TABLE ---")
        print(txt)
        print("--- END TABLE ---")
    else:
        prefix = f"[{style}] " if style else ""
        if txt.strip():
            print(prefix + txt)
